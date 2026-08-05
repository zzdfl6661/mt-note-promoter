# -*- coding: utf-8 -*-
"""生成推广笔记统计报表: 按门店分类 xlsx + csv"""
import sqlite3, os, csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(ROOT, 'data', 'promoted.db')
OUT_DIR = os.path.join(ROOT, 'reports')
os.makedirs(OUT_DIR, exist_ok=True)
stamp = datetime.now().strftime('%Y%m%d')

c = sqlite3.connect(DB)

# ---- 按门店汇总 ----
rows = c.execute("""
    SELECT store,
           COUNT(*) AS cnt,
           ROUND(AVG(view_count), 0) AS avg_views,
           MAX(view_count) AS max_views,
           MAX(promoted_at) AS last_at
    FROM promoted_notes
    GROUP BY store
    ORDER BY cnt DESC
""").fetchall()

total = sum(r[1] for r in rows)
total_avg = c.execute("SELECT ROUND(AVG(view_count),0) FROM promoted_notes").fetchone()[0]
total_max = c.execute("SELECT MAX(view_count) FROM promoted_notes").fetchone()[0]
last_global = c.execute("SELECT MAX(promoted_at) FROM promoted_notes").fetchone()[0]

def brand_of(store):
    if '鬼十八' in store:
        return '鬼十八'
    if 'bb' in store:
        return 'bb boom'
    if 'Xcape' in store or '异时刻' in store:
        return 'Xcape异时刻'
    return '其他'

# ---- 明细 ----
dets = c.execute("""
    SELECT id, store, view_count, promoted_at, substr(title, 1, 40)
    FROM promoted_notes ORDER BY store, promoted_at
""").fetchall()

# ============ XLSX ============
wb = Workbook()
ws = wb.active
ws.title = '按门店汇总'

header_fill = PatternFill('solid', fgColor='2F5597')
header_font = Font(bold=True, color='FFFFFF', size=11)
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center')

headers = ['序号', '品牌', '门店', '笔记数(篇)', '占比', '平均浏览量', '最高浏览量', '最近推广时间']
ws.append(headers)
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=i)
    cell.fill = header_fill; cell.font = header_font
    cell.alignment = center; cell.border = border

for idx, (store, cnt, avgv, maxv, last) in enumerate(rows, 1):
    ws.append([idx, brand_of(store), store, cnt,
               f'{cnt/total*100:.1f}%', int(avgv or 0), maxv or 0, (last or '')[:19]])
    for col in range(1, 9):
        ws.cell(row=idx + 1, column=col).border = border
        ws.cell(row=idx + 1, column=col).alignment = Alignment(vertical='center')

# 总计行
tr = len(rows) + 2
ws.append(['', '', '合计', total, '100%', int(total_avg or 0), total_max or 0, (last_global or '')[:19]])
for col in range(1, 9):
    cell = ws.cell(row=tr, column=col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill('solid', fgColor='FFF2CC')
    cell.border = border

widths = [6, 12, 52, 10, 8, 12, 12, 20]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'A2'

# ---- 明细 sheet ----
ws2 = wb.create_sheet('全部笔记明细')
h2 = ['ID', '门店', '浏览量', '推广时间', '笔记标题(前40字)']
ws2.append(h2)
for i, h in enumerate(h2, 1):
    cell = ws2.cell(row=1, column=i)
    cell.fill = header_fill; cell.font = header_font
    cell.alignment = center; cell.border = border
for d in dets:
    ws2.append([d[0], d[1], d[2], (d[3] or '')[:19], d[4]])
    r = ws2.max_row
    for col in range(1, 6):
        ws2.cell(row=r, column=col).border = border
for i, w in enumerate([8, 52, 10, 20, 50], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = 'A2'

xlsx_path = os.path.join(OUT_DIR, f'推广笔记统计_{stamp}.xlsx')
wb.save(xlsx_path)

# ============ CSV ============
csv_path = os.path.join(OUT_DIR, f'推广笔记统计_{stamp}.csv')
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['品牌', '门店', '笔记数(篇)', '占比', '平均浏览量', '最高浏览量', '最近推广时间'])
    for store, cnt, avgv, maxv, last in rows:
        w.writerow([brand_of(store), store, cnt, f'{cnt/total*100:.1f}%', int(avgv or 0), maxv or 0, (last or '')[:19]])
    w.writerow(['', '合计', total, '100%', int(total_avg or 0), total_max or 0, (last_global or '')[:19]])

print(f'XLSX: {xlsx_path}')
print(f'CSV : {csv_path}')
print(f'总计 {total} 篇 / {len(rows)} 家门店')
